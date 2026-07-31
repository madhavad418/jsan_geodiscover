"""A city with no comparison run must report gaps, never borrow another city's numbers."""

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import Project
from app.services.city_analytics import (
    BENCHMARK_STUDY,
    CITY_WORKBOOK,
    PORTFOLIO,
    dataset,
)
from app.services.customer_pack import build_customer_pack
from app.services.validation_layer import append_rows, get_or_create_city


def _row(place_id, name, address):
    return (
        "GRAB", place_id, name, "hair salon", address, "Bali", None, "-8.65", "115.21",
        None, None, None, "PROBE", "STRONG", None, "5", "41.5", None, "reason", "NO",
    ) + (None,) * 15


@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    session.add(Project(id=1, name="T", customer="C", country="ID"))
    session.commit()
    yield session
    session.close()


def _seed_sanur(db):
    city = get_or_create_city(db, 1, "Sanur", "Bali")
    append_rows(db, 1, city, [_row("S1", "Peni Weni", "Jl. Sudamala"), _row("S2", "Glo Spa", "Jl. Danau Poso")], "sanur.xlsx")
    db.commit()
    return city


def test_bundled_study_backs_the_benchmarked_city(db):
    city = get_or_create_city(db, 1, "Denpasar", "Bali")
    db.commit()
    data = dataset(db, 1, city.id)
    assert data["basis"] == BENCHMARK_STUDY
    assert data["comparison_available"] is True
    assert data["study"]["headline"]["total_matched_pairs"] == 53


def test_uploaded_city_reports_no_comparison_run(db):
    city = _seed_sanur(db)
    data = dataset(db, 1, city.id)
    headline = data["study"]["headline"]
    assert data["basis"] == CITY_WORKBOOK
    assert data["comparison_available"] is False
    # Never borrow the benchmarked city's pairs.
    assert headline["total_matched_pairs"] is None
    assert headline["possible_pairings"] is None
    assert headline["records_total"] == 2
    assert headline["tier1_direct_probe_candidates"] == 2
    assert data["matches"] == []


def test_city_filter_excludes_other_cities(db):
    sanur = _seed_sanur(db)
    ubud = get_or_create_city(db, 1, "Ubud", "Bali")
    append_rows(db, 1, ubud, [_row("U1", "Milk Salon", "Jl. Raya Ubud")], "ubud.xlsx")
    db.commit()
    names = {c["business_name"] for c in dataset(db, 1, sanur.id)["candidates"]}
    assert names == {"Peni Weni", "Glo Spa"}
    assert {c["business_name"] for c in dataset(db, 1, ubud.id)["candidates"]} == {"Milk Salon"}


def test_portfolio_sums_every_city(db):
    get_or_create_city(db, 1, "Denpasar", "Bali")
    _seed_sanur(db)
    db.commit()
    data = dataset(db, 1, None)
    assert data["basis"] == PORTFOLIO
    # 968 study candidates + the 2 Sanur workbook rows.
    assert data["study"]["headline"]["records_total"] == 970
    assert data["study"]["headline"]["grab_inventory"] == 402


def test_unknown_city_is_rejected(db):
    with pytest.raises(ValueError):
        dataset(db, 1, 999)


def test_generated_pack_states_the_missing_comparison(db):
    city = _seed_sanur(db)
    stream, filename = build_customer_pack(dataset(db, 1, city.id))
    assert filename == "GeoDiscover_Customer_Pack_Sanur.xlsx"

    import openpyxl
    book = openpyxl.load_workbook(stream)
    assert book.sheetnames == ["Executive Summary", "Priority Candidates", "Work Queue", "Method & Limitations"]
    summary = dict(
        (row[0], row[1]) for row in book["Executive Summary"].iter_rows(min_row=2, max_col=2, values_only=True)
    )
    assert summary["Matched pairs"] == "No comparison run for this city"
    assert summary["Businesses loaded"] == 2
    assert book["Work Queue"].max_row - 1 == 2
