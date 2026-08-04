"""Each city's workbook is appended to the shared tables without colliding."""

import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import Project, ValidationRecord
from app.services.validation_layer import append_rows, get_or_create_city, slugify


def _row(place_id, name, address):
    """A workbook row in the aligned (probed) shape."""
    return (
        "GRAB", place_id, name, "hair salon", address, "Bali", None, "-8.65", "115.21",
        None, None, None, "PROBE", "STRONG", None, "5", "41.5", None, "reason", "NO",
    ) + (None,) * 15


@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    session.add(Project(id=1, name="T", customer="C", country="ID"))
    session.commit()
    yield session
    session.close()


def test_slugify_folds_case_and_punctuation():
    assert slugify("Denpasar") == "denpasar"
    assert slugify("  Kuta Utara ") == "kuta-utara"
    assert slugify("Nusa Dua!") == "nusa-dua"


def test_same_city_name_is_never_created_twice(db):
    first = get_or_create_city(db, 1, "Ubud", "Bali")
    again = get_or_create_city(db, 1, "  ubud  ")
    assert first.id == again.id


def test_two_cities_keep_their_own_records(db):
    ubud = get_or_create_city(db, 1, "Ubud")
    sanur = get_or_create_city(db, 1, "Sanur")
    append_rows(db, 1, ubud, [_row("U1", "Milk Salon", "Jl. Raya Ubud")], "ubud.xlsx")
    append_rows(db, 1, sanur, [_row("S1", "Peni Weni", "Jl. Sudamala")], "sanur.xlsx")
    db.commit()
    assert db.scalar(select(ValidationRecord).where(ValidationRecord.city_id == ubud.id)).business_name == "Milk Salon"
    assert len(list(db.scalars(select(ValidationRecord)))) == 2


def test_reupload_refreshes_instead_of_duplicating(db):
    city = get_or_create_city(db, 1, "Ubud")
    rows = [_row("U1", "Milk Salon", "Jl. Raya Ubud"), _row("U2", "Bloom", "Jl. Monkey Forest")]
    first = append_rows(db, 1, city, rows, "ubud.xlsx")
    second = append_rows(db, 1, city, rows, "ubud-v2.xlsx")
    db.commit()
    assert (first["added"], second["added"]) == (2, 0)
    assert second["refreshed"] == 2
    assert len(list(db.scalars(select(ValidationRecord)))) == 2
    assert city.source_files == ["ubud.xlsx", "ubud-v2.xlsx"]


def test_reupload_keeps_saved_validation_answers(db):
    city = get_or_create_city(db, 1, "Ubud")
    append_rows(db, 1, city, [_row("U1", "Milk Salon", "Jl. Raya Ubud")], "ubud.xlsx")
    db.commit()
    record = db.scalar(select(ValidationRecord))
    record.fresha = "YES"
    record.fresha_link = "https://www.fresha.com/a/milk"
    record.validation_status = "IN_PROGRESS"
    db.commit()

    append_rows(db, 1, city, [_row("U1", "Milk Salon Renamed", "Jl. Raya Ubud 2")], "ubud.xlsx")
    db.commit()
    db.refresh(record)
    assert record.business_name == "Milk Salon Renamed"   # discovery columns refresh
    assert record.fresha == "YES"                          # saved answers survive
    assert record.fresha_link == "https://www.fresha.com/a/milk"


def test_rows_without_a_place_id_dedupe_on_name_and_address(db):
    city = get_or_create_city(db, 1, "Ubud")
    row = _row(None, "Milk Salon", "Jl. Raya Ubud")
    append_rows(db, 1, city, [row], "a.xlsx")
    second = append_rows(db, 1, city, [row], "b.xlsx")
    db.commit()
    assert second["added"] == 0
    assert len(list(db.scalars(select(ValidationRecord)))) == 1


def test_city_validation_export_contains_current_database_values(client, headers):
    project_response = client.post("/api/v1/projects", headers=headers, json={
        "name": "City export test",
        "customer": "Demo Customer",
        "country": "Indonesia",
        "cities": ["Ubud"],
        "categories": ["Hair Salon"],
        "target_volume": 10,
        "confidence_threshold": 70,
        "gold_threshold": 85,
        "require_qa_for_release": True,
    })
    assert project_response.status_code == 201, project_response.text
    project_id = project_response.json()["id"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "Provider", "Provider Place ID", "Business Name", "Categories", "Address",
        "City", "Province", "Latitude", "Longitude", "Evidence Method", "Google maps",
    ])
    sheet.append(["GRAB", "U1", "Milk Salon", "hair salon", "Jl. Raya Ubud", "Ubud", "Bali", -8.5, 115.2, "PROBE", None])
    upload = io.BytesIO()
    workbook.save(upload)

    uploaded = client.post(
        f"/api/v1/cities/{project_id}/upload",
        headers=headers,
        data={"city": "Ubud", "region": "Bali", "country": "Indonesia"},
        files={"file": ("ubud.xlsx", upload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert uploaded.status_code == 200, uploaded.text
    city_id = uploaded.json()["city"]["id"]

    records = client.get(
        f"/api/v1/validation/{project_id}/records?city_id={city_id}", headers=headers
    ).json()["records"]
    saved = client.put(
        f"/api/v1/validation/records/{records[0]['id']}",
        headers=headers,
        json={
            "google_maps": "YES",
            "google_maps_link": "https://maps.google.com/milk",
            "validation_notes": "Checked by analyst",
        },
    )
    assert saved.status_code == 200, saved.text

    exported = client.get(
        f"/api/v1/validation/{project_id}/cities/{city_id}/xlsx", headers=headers
    )
    assert exported.status_code == 200, exported.text
    assert "GeoDiscover_Ubud_Validation_" in exported.headers["content-disposition"]
    result = load_workbook(io.BytesIO(exported.content), data_only=True)
    rows = list(result["Validation Data"].iter_rows(values_only=True))
    values = dict(zip(rows[0], rows[1]))
    assert values["Business Name"] == "Milk Salon"
    assert values["Google Maps"] == "YES"
    assert values["Google Maps Link"] == "https://maps.google.com/milk"
    assert values["Validator notes"] == "Checked by analyst"
