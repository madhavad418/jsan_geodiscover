import csv
import io
import uuid
from datetime import date
from typing import Annotated, Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..auth import require_roles
from ..database import get_db
from ..models import CandidateStatus, GoogleStatus, PoiCandidate, Project, RawSourceRecord, User, UserRole
from ..schemas import ImportResult
from ..services.audit import write_audit
from ..services.eligibility import recalculate_candidate
from ..services.normalization import normalize_row

router = APIRouter(prefix="/imports", tags=["Imports"])


def _read_rows(file_name: str, content: bytes) -> list[dict[str, Any]]:
    lower = file_name.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    if lower.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(v).strip() if v is not None else "" for v in values[0]]
        return [dict(zip(headers, row)) for row in values[1:] if any(v not in (None, "") for v in row)]
    raise HTTPException(400, "Only CSV and XLSX files are supported")


@router.post("/poi", response_model=ImportResult)
async def import_poi_file(
    project_id: Annotated[int, Form()],
    source_name: Annotated[str, Form()],
    file: Annotated[UploadFile, File()],
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(require_roles(UserRole.PROGRAM_ADMIN, UserRole.PROJECT_MANAGER, UserRole.DISCOVERY_ANALYST))],
):
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(404, "Project not found")
    content = await file.read()
    rows = _read_rows(file.filename or "upload.csv", content)
    imported_ids: list[int] = []
    rejected = 0
    for index, row in enumerate(rows, start=2):
        normalized = normalize_row(row)
        if normalized["business_name"] == "Unnamed candidate":
            rejected += 1
            continue
        raw = RawSourceRecord(
            project_id=project.id,
            source_name=source_name,
            source_record_id=str(normalized.get("source_record_id") or "") or None,
            raw_payload={str(k): v for k, v in row.items()},
            file_name=file.filename,
            imported_by=user.id,
        )
        db.add(raw)
        db.flush()
        candidate = PoiCandidate(
            project_id=project.id,
            raw_source_record_id=raw.id,
            poi_id=f"POI-{uuid.uuid4().hex[:12].upper()}",
            source_name=source_name,
            source_type="FILE_UPLOAD",
            source_record_id=normalized.pop("source_record_id", None),
            coordinate_precision="ROOFTOP" if normalized.get("latitude") is not None and normalized.get("longitude") is not None else "UNKNOWN",
            extraction_date=date.today(),
            reviewer_identity=f"import:{user.email}",
            created_by=user.id,
            updated_by=user.id,
            confidence_score=55,
            google_status=GoogleStatus.GOOGLE_CHECK_PENDING,
            google_review_completed=False,
            status=CandidateStatus.GOOGLE_CHECK_PENDING,
            **normalized,
        )
        recalculate_candidate(candidate, project)
        db.add(candidate)
        db.flush()
        imported_ids.append(candidate.id)
    write_audit(db, user, "POI_FILE_IMPORTED", "project", project.id, {
        "file_name": file.filename,
        "source_name": source_name,
        "imported": len(imported_ids),
        "rejected_rows": rejected,
    })
    db.commit()
    return ImportResult(imported=len(imported_ids), rejected_rows=rejected, candidate_ids=imported_ids)
