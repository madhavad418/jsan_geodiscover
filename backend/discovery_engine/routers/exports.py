import io
from datetime import datetime, timezone
from typing import Annotated
import re
import unicodedata
from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..auth import require_roles
from ..database import get_db
from ..models import ExportRecord, PoiCandidate, Project, User, UserRole
from ..services.audit import write_audit
from ..services.kpi import calculate_kpis

router = APIRouter(prefix="/exports", tags=["Exports"])

HEADERS = [
    "POI ID", "Business Name", "Category", "Address", "City", "Province", "Latitude", "Longitude",
    "Coordinate Precision", "Phone", "WhatsApp", "Website", "Booking URL", "Instagram", "TikTok", "Facebook",
    "Source", "Source Type", "Source Record ID", "Primary Source URL", "Evidence Reference", "Extraction Date",
    "Google Status", "Confidence", "Quality Tier", "Usable Non-Google", "Validation Date", "Demo Record"
]


@router.get("/{project_id}/xlsx")
def export_xlsx(
    project_id: int,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(require_roles(UserRole.PROGRAM_ADMIN, UserRole.PROJECT_MANAGER, UserRole.QA_REVIEWER, UserRole.CUSTOMER_VIEWER))],
    usable_only: bool = True,
):
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(404, "Project not found")
    query = select(PoiCandidate).where(PoiCandidate.project_id == project_id)
    if usable_only:
        query = query.where(PoiCandidate.usability_eligible.is_(True))
    rows = db.scalars(query.order_by(PoiCandidate.city, PoiCandidate.business_name)).all()
    all_rows = db.scalars(select(PoiCandidate).where(PoiCandidate.project_id == project_id)).all()
    kpis = calculate_kpis(all_rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "POI Deliverable"
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183B56")
    for r in rows:
        ws.append([
            r.poi_id, r.business_name, r.category, r.full_address, r.city, r.province, r.latitude, r.longitude,
            r.coordinate_precision, r.phone, r.whatsapp, r.website, r.booking_url, r.instagram, r.tiktok, r.facebook,
            r.source_name, r.source_type, r.source_record_id, r.primary_source_url, r.evidence_reference,
            r.extraction_date.isoformat() if r.extraction_date else None, r.google_status.value, r.confidence_score,
            r.quality_tier.value, "YES" if r.usability_eligible else "NO", r.updated_at.isoformat(), "YES" if r.is_demo else "NO"
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(max(len(str(c.value or "")) for c in column) + 2, 42)

    summary = wb.create_sheet("Validation Summary")
    summary.append(["Metric", "Value"])
    for key, value in kpis.items():
        if not isinstance(value, list):
            summary.append([key, value])
    summary["A1"].font = summary["B1"].font = Font(bold=True, color="FFFFFF")
    summary["A1"].fill = summary["B1"].fill = PatternFill("solid", fgColor="183B56")

    buffer = io.BytesIO()
    wb.save(buffer)
    safe_name = unicodedata.normalize("NFKD", project.name).encode("ascii", "ignore").decode("ascii")
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", safe_name).strip("_") or f"project_{project.id}"
    file_name = f"GeoDiscover_{safe_name}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
    db.add(ExportRecord(project_id=project_id, file_name=file_name, record_count=len(rows), validation_summary=kpis, created_by=user.id))
    write_audit(db, user, "CUSTOMER_EXPORT_CREATED", "project", project_id, {"file_name": file_name, "record_count": len(rows), "usable_only": usable_only})
    db.commit()
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )
