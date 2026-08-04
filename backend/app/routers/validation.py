from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Annotated, Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..auth import get_current_user
from ..database import get_db
from ..models import City, Project, User, ValidationRecord
from ..schemas import ValidationRecordUpdate
from ..services.audit import write_audit
from ..services.validation_layer import (
    CHECK_FIELDS,
    CHECK_KEYS,
    EDITABLE_KEYS,
    LINK_FIELDS,
    LINK_KEYS_BY_CHECK,
    SOURCE_FIELDS,
    compute_status,
    field_schema,
    import_validation_folder,
    normalise_check,
    normalise_link,
)

router = APIRouter(prefix="/validation", tags=["Validation Layer"])

LINK_KEYS = set(LINK_KEYS_BY_CHECK.values()) | {field["key"] for field in LINK_FIELDS}
CHECK_LABELS = {field["key"]: field["label"] for field in CHECK_FIELDS}


def _export_columns() -> list[tuple[str, str]]:
    columns = [(field["label"], field["key"]) for field in SOURCE_FIELDS]
    for field in CHECK_FIELDS:
        columns.extend([(field["label"], field["key"]), (f'{field["label"]} Link', field["link_key"])])
    columns.extend((field["label"], field["key"]) for field in LINK_FIELDS)
    columns.extend((field["label"], field["key"]) for field in field_schema()["text_fields"])
    columns.extend([
        ("Validation Status", "validation_status"),
        ("Updated By", "updated_by_email"),
        ("Updated At", "updated_at"),
        ("Source File", "source_file"),
    ])
    return columns


def _require_project(project_id: int, db: Session) -> Project:
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(404, "Project not found")
    return project


def _serialise(record: ValidationRecord) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "id": record.id,
        "source_row": record.source_row,
        "source_file": record.source_file,
        "validation_status": record.validation_status,
        "updated_by_email": record.updated_by_email,
        "updated_at": record.updated_at.isoformat() if record.updated_at else None,
    }
    for field in SOURCE_FIELDS:
        payload[field["key"]] = getattr(record, field["key"])
    for key in EDITABLE_KEYS:
        payload[key] = getattr(record, key)
    return payload


@router.get("/schema")
def validation_schema(user: Annotated[User, Depends(get_current_user)]) -> dict[str, Any]:
    return field_schema()


@router.get("/{project_id}/records")
def list_records(
    project_id: int,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(get_current_user)],
    search: str | None = Query(None),
    provider: str | None = Query(None),
    status: str | None = Query(None),
    city_id: int | None = Query(None, description="Restrict to one city; omit for every city"),
) -> dict[str, Any]:
    _require_project(project_id, db)

    # First open of the tab on a database seeded before this feature existed.
    scope = ValidationRecord.project_id == project_id
    if not db.scalar(select(func.count(ValidationRecord.id)).where(scope)):
        import_validation_folder(db, project_id)
    if city_id:
        scope = scope & (ValidationRecord.city_id == city_id)

    query = select(ValidationRecord).where(scope)
    if provider:
        query = query.where(ValidationRecord.provider == provider.upper())
    if status:
        query = query.where(ValidationRecord.validation_status == status.upper())
    if search:
        term = f"%{search.strip()}%"
        query = query.where(
            ValidationRecord.business_name.ilike(term) | ValidationRecord.address.ilike(term)
        )

    records = list(db.scalars(query.order_by(ValidationRecord.city_id, ValidationRecord.source_row)))
    all_records = list(db.scalars(select(ValidationRecord).where(scope)))
    counts = {"PENDING": 0, "IN_PROGRESS": 0, "COMPLETED": 0}
    for record in all_records:
        counts[record.validation_status] = counts.get(record.validation_status, 0) + 1

    cities = list(db.scalars(select(City).where(City.project_id == project_id).order_by(City.name)))
    city_names = {city.id: city.name for city in cities}
    city_counts = dict(
        db.execute(
            select(ValidationRecord.city_id, func.count())
            .where(ValidationRecord.project_id == project_id)
            .group_by(ValidationRecord.city_id)
        ).all()
    )

    return {
        "project_id": project_id,
        "city_id": city_id,
        "cities": [
            {"id": c.id, "name": c.name, "region": c.region, "records": city_counts.get(c.id, 0)}
            for c in cities
        ],
        "total": len(all_records),
        "returned": len(records),
        "summary": counts,
        "providers": sorted({r.provider for r in all_records if r.provider}),
        "source_file": next((r.source_file for r in all_records if r.source_file), None),
        "records": [{**_serialise(r), "city": city_names.get(r.city_id)} for r in records],
    }


@router.get("/{project_id}/cities/{city_id}/xlsx")
def export_city_validation_workbook(
    project_id: int,
    city_id: int,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(get_current_user)],
) -> Response:
    """Export one city's current validation records directly from the database."""
    project = _require_project(project_id, db)
    city = db.get(City, city_id)
    if not city or city.project_id != project_id:
        raise HTTPException(404, "City not found")

    records = list(db.scalars(
        select(ValidationRecord)
        .where(ValidationRecord.project_id == project_id, ValidationRecord.city_id == city_id)
        .order_by(ValidationRecord.source_row)
    ))
    if not records:
        raise HTTPException(422, "This city has no validation records to export")

    columns = _export_columns()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Validation Data"
    sheet.append([label for label, _ in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183B56")
    for record in records:
        values = []
        for _, key in columns:
            value = getattr(record, key)
            if key == "updated_at" and value is not None:
                value = value.isoformat()
            values.append(value)
        sheet.append(values)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)

    info = workbook.create_sheet("Export Info")
    info.append(["Field", "Value"])
    info.append(["Project", project.name])
    info.append(["City", city.name])
    info.append(["Region", city.region])
    info.append(["Country", city.country])
    info.append(["Record Count", len(records)])
    info.append(["Generated By", user.email])
    info.append(["Generated At", datetime.now(timezone.utc).isoformat()])
    for cell in info[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183B56")
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 45

    output = io.BytesIO()
    workbook.save(output)
    safe_city = re.sub(r"[^A-Za-z0-9_-]+", "_", city.name).strip("_") or f"city_{city.id}"
    file_name = f"GeoDiscover_{safe_city}_Validation_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
    write_audit(db, user, "CITY_VALIDATION_EXPORTED", "City", city.id, {"file_name": file_name, "record_count": len(records)})
    db.commit()
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )


@router.put("/records/{record_id}")
def save_record(
    record_id: int,
    payload: ValidationRecordUpdate,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(get_current_user)],
) -> dict[str, Any]:
    record = db.get(ValidationRecord, record_id)
    if not record:
        raise HTTPException(404, "Validation record not found")

    submitted = payload.model_dump()

    # Normalise the YES/NO answers first so the link rules can rely on them.
    answers: dict[str, str | None] = {}
    for key in CHECK_KEYS:
        raw = submitted.get(key)
        answer = normalise_check(raw)
        if raw not in (None, "") and answer is None:
            raise HTTPException(422, f"{CHECK_LABELS[key]} must be answered Yes or No")
        answers[key] = answer

    values: dict[str, Any] = {}
    for key in EDITABLE_KEYS:
        if key in CHECK_KEYS:
            values[key] = answers[key]
        elif key in LINK_KEYS:
            values[key] = normalise_link(submitted.get(key))
        else:
            text = str(submitted.get(key) or "").strip()
            values[key] = text or None

    # A YES answer has to be backed by a link; a NO answer never carries one.
    for key in CHECK_KEYS:
        link_key = LINK_KEYS_BY_CHECK[key]
        if answers[key] == "YES" and not values[link_key]:
            raise HTTPException(422, f"{CHECK_LABELS[key]} is marked Yes - add the supporting link")
        if answers[key] != "YES":
            values[link_key] = None

    before = {key: getattr(record, key) for key in EDITABLE_KEYS}
    for key, value in values.items():
        setattr(record, key, value)
    record.validation_status = compute_status(record)
    record.updated_by_email = user.email

    changed = {key: {"from": before[key], "to": values[key]} for key in EDITABLE_KEYS if before[key] != values[key]}
    write_audit(
        db,
        user,
        "VALIDATION_RECORD_SAVED",
        "ValidationRecord",
        record.id,
        {"business_name": record.business_name, "status": record.validation_status, "changed": changed},
    )
    db.commit()
    db.refresh(record)
    return _serialise(record)


@router.post("/{project_id}/reimport")
def reimport(
    project_id: int,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(get_current_user)],
) -> dict[str, Any]:
    """Refresh the discovery columns from every workbook. Saved answers are kept."""
    _require_project(project_id, db)
    results = import_validation_folder(db, project_id)
    if not results:
        raise HTTPException(404, "No .xlsx workbook found in backend/misseddata")
    write_audit(db, user, "VALIDATION_WORKBOOK_IMPORTED", "Project", project_id, {"results": results})
    db.commit()
    return {"imported": results}
