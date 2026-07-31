import io
import re
import unicodedata
from datetime import datetime, timezone
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..auth import require_roles
from ..database import get_db
from ..models import EvidenceItem, ExportRecord, GoogleCheck, PoiCandidate, Project, QaStatus, User, UserRole
from ..services.audit import write_audit
from ..services.kpi import calculate_kpis

router = APIRouter(prefix="/exports", tags=["Exports"])

HEADERS = [
    "POI ID", "Business Name", "Local Name", "Alternate Names", "Category", "Subcategory", "Branch",
    "Address", "Locality", "District", "City", "Province", "Postal Code", "Latitude", "Longitude",
    "Coordinate Precision", "Phone", "WhatsApp", "WhatsApp Business", "Email", "Website", "Booking URL",
    "Instagram", "TikTok", "Facebook", "Source", "Source Type", "Source Record ID", "Primary Source URL",
    "Evidence Reference", "Extraction Date", "Last Active Date", "Google Status", "Comparison Bucket",
    "Confidence", "Confidence Components", "Quality Tier", "Usable Non-Google", "QA Status", "QA Reviewer",
    "Release Approved", "Validation Date", "Demo Record"
]


def _project_rows(db: Session, project_id: int) -> list[PoiCandidate]:
    return list(db.scalars(select(PoiCandidate).where(PoiCandidate.project_id == project_id)).all())


@router.get("/{project_id}/validation")
def validate_export(
    project_id: int,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(require_roles(UserRole.PROGRAM_ADMIN, UserRole.PROJECT_MANAGER, UserRole.QA_REVIEWER, UserRole.CUSTOMER_VIEWER))],
):
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(404, "Project not found")
    all_rows = _project_rows(db, project_id)
    release_rows = [row for row in all_rows if row.release_approved and not row.is_demo]
    blocked = [{
        "candidate_id": row.id,
        "poi_id": row.poi_id,
        "business_name": row.business_name,
        "qa_status": row.qa_status.value,
        "is_demo": row.is_demo,
        "usability_failures": row.usability_failures,
    } for row in all_rows if row.usability_eligible and not row.release_approved]
    return {
        "project_id": project_id,
        "customer": project.customer,
        "release_ready_records": len(release_rows),
        "blocked_usable_records": blocked,
        "demo_records_excluded": sum(row.is_demo for row in all_rows),
        "qa_passed": sum(row.qa_status == QaStatus.PASSED for row in all_rows),
        "kpis": calculate_kpis(all_rows),
        "customer_export_allowed": bool(release_rows),
    }


@router.get("/{project_id}/xlsx")
def export_xlsx(
    project_id: int,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(require_roles(UserRole.PROGRAM_ADMIN, UserRole.PROJECT_MANAGER, UserRole.QA_REVIEWER, UserRole.CUSTOMER_VIEWER))],
    release_only: bool = True,
    include_demo: bool = False,
):
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(404, "Project not found")
    query = select(PoiCandidate).where(PoiCandidate.project_id == project_id)
    if release_only:
        query = query.where(PoiCandidate.release_approved.is_(True))
    if not include_demo:
        query = query.where(PoiCandidate.is_demo.is_(False))
    rows = list(db.scalars(query.order_by(PoiCandidate.city, PoiCandidate.business_name)).all())
    all_rows = _project_rows(db, project_id)
    kpis = calculate_kpis(all_rows)
    if release_only and not rows:
        raise HTTPException(422, "No customer-release records are available. Records require D eligibility, independent QA pass, and non-demo status.")

    wb = Workbook()
    ws = wb.active
    ws.title = "POI Deliverable"
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183B56")
    for r in rows:
        ws.append([
            r.poi_id, r.business_name, r.local_name, "; ".join(r.alternate_names or []), r.category, r.subcategory,
            r.branch_name, r.full_address, r.locality, r.district, r.city, r.province, r.postal_code, r.latitude,
            r.longitude, r.coordinate_precision, r.phone, r.whatsapp, "YES" if r.whatsapp_business else "NO",
            r.email, r.website, r.booking_url, r.instagram, r.tiktok, r.facebook, r.source_name, r.source_type,
            r.source_record_id, r.primary_source_url, r.evidence_reference,
            r.extraction_date.isoformat() if r.extraction_date else None,
            r.last_active_date.isoformat() if r.last_active_date else None,
            r.google_status.value, r.comparison_bucket.value, r.confidence_score, str(r.confidence_components or {}),
            r.quality_tier.value, "YES" if r.usability_eligible else "NO", r.qa_status.value, r.qa_reviewer,
            "YES" if r.release_approved else "NO", r.updated_at.isoformat(), "YES" if r.is_demo else "NO",
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(max(len(str(c.value or "")) for c in column) + 2, 45)

    summary = wb.create_sheet("Validation Summary")
    summary.append(["Metric", "Value"])
    for key, value in kpis.items():
        if not isinstance(value, list):
            summary.append([key, value])
    summary.append(["export_release_only", release_only])
    summary.append(["export_include_demo", include_demo])
    summary.append(["export_record_count", len(rows)])
    summary.append(["generated_by", user.email])
    summary.append(["generated_at", datetime.now(timezone.utc).isoformat()])
    summary["A1"].font = summary["B1"].font = Font(bold=True, color="FFFFFF")
    summary["A1"].fill = summary["B1"].fill = PatternFill("solid", fgColor="183B56")

    evidence_ws = wb.create_sheet("Evidence Manifest")
    evidence_ws.append(["POI ID", "Evidence ID", "Type", "Title", "Source URL", "File Name", "SHA-256", "Created At"])
    for r in rows:
        for item in db.scalars(select(EvidenceItem).where(EvidenceItem.poi_candidate_id == r.id)).all():
            evidence_ws.append([r.poi_id, item.id, item.evidence_type.value, item.title, item.source_url, item.file_name, item.sha256, item.created_at.isoformat()])

    google_ws = wb.create_sheet("Google Exclusion Audit")
    google_ws.append(["POI ID", "Decision", "Reviewer", "Mandatory Checks Complete", "Check Type", "Query", "Result", "Place IDs", "Evidence", "Timestamp"])
    for r in rows:
        checks = db.scalars(select(GoogleCheck).where(GoogleCheck.poi_candidate_id == r.id).order_by(GoogleCheck.created_at)).all()
        for check in checks:
            if check.checklist:
                for item in check.checklist:
                    google_ws.append([r.poi_id, check.decision.value, check.reviewer, check.mandatory_checks_complete, item.get("check_type"), item.get("query"), item.get("result"), ", ".join(item.get("candidate_place_ids") or []), item.get("evidence_reference") or check.evidence_reference, check.created_at.isoformat()])
            else:
                google_ws.append([r.poi_id, check.decision.value, check.reviewer, check.mandatory_checks_complete, None, None, None, ", ".join(check.candidate_place_ids or []), check.evidence_reference, check.created_at.isoformat()])

    buffer = io.BytesIO()
    wb.save(buffer)
    safe_name = unicodedata.normalize("NFKD", project.name).encode("ascii", "ignore").decode("ascii")
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", safe_name).strip("_") or f"project_{project.id}"
    file_name = f"GeoDiscover_{safe_name}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
    db.add(ExportRecord(project_id=project_id, file_name=file_name, record_count=len(rows), validation_summary=kpis, created_by=user.id))
    write_audit(db, user, "CUSTOMER_EXPORT_CREATED", "project", project_id, {"file_name": file_name, "record_count": len(rows), "release_only": release_only, "include_demo": include_demo})
    db.commit()
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )
