"""Build a customer workbook for a city that has no signed-off pack of its own."""

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="102F46")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13)


def _sheet(book: Workbook, title: str, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None):
    sheet = book.create_sheet(title[:31])
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(row)
    for index, width in enumerate(widths or [], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    return sheet


def build_customer_pack(data: dict[str, Any]) -> tuple[io.BytesIO, str]:
    """Return an in-memory workbook and its filename for the given city dataset."""
    study = data["study"]
    headline = study["headline"]
    candidates = data["candidates"]
    where = study["geography"]

    book = Workbook()
    book.remove(book.active)

    # 1. Executive summary
    summary = book.create_sheet("Executive Summary")
    summary["A1"] = f"JSAN GeoDiscover - {study['title']}"
    summary["A1"].font = TITLE_FONT
    rows = [
        ("Customer", study["customer"]),
        ("Geography", where),
        ("Data basis", data["basis_label"]),
        ("Source workbook", study["source_workbook"]),
        ("", ""),
        ("Businesses loaded", headline.get("records_total")),
        ("Grab inventory", headline.get("grab_inventory")),
        ("Google inventory", headline.get("google_inventory")),
        ("Matched pairs", headline.get("total_matched_pairs") if headline.get("total_matched_pairs") is not None else "No comparison run for this city"),
        ("Tier 1 direct-probe candidates", headline.get("tier1_direct_probe_candidates")),
        ("Tier 2 area-supported leads", headline.get("tier2_area_supported_leads")),
        ("Contactability %", headline.get("contactability_pct")),
        ("Validation completed", headline.get("validation_completed")),
        ("Validation in progress", headline.get("validation_in_progress")),
        ("Records scraped", headline.get("scraped_records")),
        ("Customer release ready", headline.get("customer_release_ready", 0)),
        ("", ""),
        ("Safe customer statement", data.get("customer_message") or ""),
        ("Do not state", "Any candidate is a confirmed active business absent from Google"),
    ]
    for label, value in rows:
        summary.append([label, value])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 78
    for row in summary.iter_rows(min_row=2):
        row[0].font = Font(bold=True, size=10)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # 2. Priority candidates
    tier1 = [c for c in candidates if c.get("customer_tier") == "TIER_1_DIRECT_PROBE"]
    _sheet(
        book,
        "Priority Candidates",
        ["Rank", "City", "Business name", "Address", "Latitude", "Longitude", "Evidence method", "Evidence reason", "Contact number", "Validation status"],
        [
            [
                index,
                row.get("city"),
                row.get("business_name"),
                row.get("address"),
                row.get("latitude"),
                row.get("longitude"),
                row.get("evidence_method"),
                row.get("evidence_reason"),
                row.get("identified_contact_number"),
                row.get("validation_status"),
            ]
            for index, row in enumerate(tier1, start=1)
        ],
        [7, 16, 34, 46, 12, 12, 18, 52, 20, 18],
    )

    # 3. Full work queue
    _sheet(
        book,
        "Work Queue",
        ["City", "Provider", "Business name", "Categories", "Address", "Latitude", "Longitude", "Evidence method", "Tier", "Contact number", "Validation status"],
        [
            [
                row.get("city"),
                row.get("provider"),
                row.get("business_name"),
                row.get("categories"),
                row.get("address"),
                row.get("latitude"),
                row.get("longitude"),
                row.get("evidence_method"),
                row.get("customer_tier") or "",
                row.get("identified_contact_number"),
                row.get("validation_status"),
            ]
            for row in candidates
        ],
        [16, 10, 34, 34, 46, 12, 12, 18, 24, 20, 18],
    )

    # 4. Method and limitations
    _sheet(
        book,
        "Method & Limitations",
        ["Item", "Statement"],
        [
            ["Data basis", data["basis_label"]],
            ["Scope", where],
            ["Comparison run", "Completed" if data.get("comparison_available") else "Not run for this city"],
            ["Note", data.get("note") or ""],
            ["Release control", "No record is customer-ready until category, contact, activity, full Google exclusion and independent QA are complete."],
            ["Empty columns", "A validation column is only filled when a provider proved a value. Empty means unproven, not absent."],
        ],
        [26, 110],
    )

    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    label = "All_Cities" if data.get("city") is None else where.split(",")[0].strip()
    slug = re.sub(r"[^A-Za-z0-9]+", "_", label) or "City"
    return stream, f"GeoDiscover_Customer_Pack_{slug}.xlsx"
